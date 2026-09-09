import os
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

CARPETA_REPORTES = "storage/reportes_generados"

def _formatear_celda_encabezado(celda, texto):
    celda.text = texto
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
    shading = parse_xml(r'<w:shd {} w:fill="1F4E79"/>'.format(nsdecls('w')))
    celda._tc.get_or_add_tcPr().append(shading)

def crear_documento_pae(paciente, valoracion, plan_pae_dict: dict) -> str:
    """Genera un archivo Word (.docx) con formato estructurado de informe PAE."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. Título Principal
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titulo = titulo.add_run("PROCESO DE ATENCIÓN DE ENFERMERÍA (PAE)")
    run_titulo.font.size = Pt(16)
    run_titulo.font.bold = True
    run_titulo.font.color.rgb = RGBColor(31, 78, 121)

    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitulo.add_run("Informe de Valoración y Plan de Cuidados")
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 2. Datos del Paciente y Valoración
    h1 = doc.add_heading("I. VALORACIÓN Y DATOS DE FILIACIÓN", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    tabla_datos = doc.add_table(rows=3, cols=2)
    tabla_datos.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla_datos.autofit = False

    datos_paciente = [
        [f"Paciente: {paciente.nombre}", f"Edad: {paciente.edad or 'N/E'} años | Sexo: {paciente.sexo or 'N/E'}"],
        [f"Cama/Servicio: {paciente.cama or 'N/E'}", f"Fecha: {paciente.fecha_registro.strftime('%d/%m/%Y')}"],
        [f"Motivo de Ingreso: {valoracion.motivo_ingreso or 'N/E'}", f"Antecedentes: {valoracion.antecedentes or 'Sin antecedentes de relevancia'}"]
    ]

    for i, fila in enumerate(datos_paciente):
        for j, texto in enumerate(fila):
            celda = tabla_datos.cell(i, j)
            celda.text = texto
            celda.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Signos vitales
    if valoracion.signos_vitales:
        try:
            sv = json.loads(valoracion.signos_vitales) if isinstance(valoracion.signos_vitales, str) else valoracion.signos_vitales
            doc.add_paragraph(f"Signos Vitales: PA: {sv.get('presion_arterial', '-')} | FC: {sv.get('frecuencia_cardiaca', '-')} | FR: {sv.get('frecuencia_respiratoria', '-')} | Temp: {sv.get('temperatura', '-')} | SatO2: {sv.get('saturacion_oxigeno', '-')}")
        except Exception:
            pass

    p_sub = doc.add_paragraph()
    p_sub.add_run("Datos Subjetivos: ").bold = True
    p_sub.add_run(valoracion.datos_subjetivos or "No especificados")

    p_obj = doc.add_paragraph()
    p_obj.add_run("Datos Objetivos: ").bold = True
    p_obj.add_run(valoracion.datos_objetivos or "No especificados")

    # 3. Plan de Cuidados (NANDA, NOC, NIC)
    h2 = doc.add_heading("II. PLANIFICACIÓN DEL CUIDADO (NANDA - NOC - NIC)", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(8)

    planes = plan_pae_dict.get("planes", [])
    for idx, plan in enumerate(planes, 1):
        diag = plan.get("diagnostico_nanda", {})
        
        p_diag = doc.add_paragraph()
        run_d = p_diag.add_run(f"Diagnóstico #{idx}: {diag.get('codigo')} - {diag.get('etiqueta')}")
        run_d.bold = True
        run_d.font.color.rgb = RGBColor(31, 78, 121)

        p_pes = doc.add_paragraph()
        p_pes.add_run("Enunciado PES: ").bold = True
        p_pes.add_run(diag.get("formato_pes", ""))

        tabla_plan = doc.add_table(rows=1, cols=3)
        tabla_plan.alignment = WD_TABLE_ALIGNMENT.CENTER

        encabezados = ["Diagnóstico NANDA", "Resultados NOC (Indicadores)", "Intervenciones NIC (Actividades)"]
        for j, h in enumerate(encabezados):
            _formatear_celda_encabezado(tabla_plan.cell(0, j), h)

        row = tabla_plan.add_row()
        
        # Columna 0: NANDA
        row.cells[0].text = f"Dominio: {diag.get('dominio', '')}\nClase: {diag.get('clase', '')}\n\n{diag.get('formato_pes', '')}"
        
        # Columna 1: NOC (aquí estaba el typo corregido)
        texto_noc = []
        for noc in plan.get("resultados_noc", []):
            texto_noc.append(f"• {noc.get('codigo')} {noc.get('nombre')}")
            for ind in noc.get("indicadores", []):
                texto_noc.append(f"  - {ind.get('nombre')} (Diana: {ind.get('puntuacion_actual')} -> {ind.get('puntuacion_diana')})")
        row.cells[1].text = "\n".join(texto_noc)

        # Columna 2: NIC
        texto_nic = []
        for nic in plan.get("intervenciones_nic", []):
            texto_nic.append(f"• {nic.get('codigo')} {nic.get('nombre')}")
            for act in nic.get("actividades", []):
                texto_nic.append(f"  - {act}")
        row.cells[2].text = "\n".join(texto_nic)

        if plan.get("evaluacion_esperada"):
            p_ev = doc.add_paragraph()
            p_ev.paragraph_format.space_before = Pt(4)
            p_ev.add_run("Evaluación esperada: ").bold = True
            p_ev.add_run(plan.get("evaluacion_esperada"))

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

# 4. Registro de Enfermería SOAPIE (Nueva sección)
    soapie_data = plan_pae_dict.get("soapie")
    if soapie_data:
        h3 = doc.add_heading("III. REGISTRO CLÍNICO DE ENFERMERÍA (SOAPIE)", level=1)
        h3.paragraph_format.space_before = Pt(14)
        h3.paragraph_format.space_after = Pt(8)

        tabla_soapie = doc.add_table(rows=1, cols=2)
        tabla_soapie.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabla_soapie.autofit = False

        _formatear_celda_encabezado(tabla_soapie.cell(0, 0), "Letra")
        _formatear_celda_encabezado(tabla_soapie.cell(0, 1), "Registro Clínico de Enfermería")
        tabla_soapie.columns[0].width = Inches(0.8)
        tabla_soapie.columns.width = Inches(5.8)

        filas_soapie = [
            ("S", "Subjetivo", soapie_data.get("s", "")),
            ("O", "Objetivo", soapie_data.get("o", "")),
            ("A", "Análisis", soapie_data.get("a", "")),
            ("P", "Planificación", soapie_data.get("p", "")),
            ("I", "Intervención", soapie_data.get("i", "")),
            ("E", "Evaluación", soapie_data.get("e", ""))
        ]

        for letra, nombre, contenido in filas_soapie:
            row = tabla_soapie.add_row()
            cell_letra = row.cells[0]
            cell_letra.text = f"{letra}\n({nombre})"
            cell_letra.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell_letra.paragraphs[0].runs[0].font.bold = True
            cell_letra.paragraphs[0].runs[0].font.size = Pt(11)
            cell_letra.paragraphs[0].runs[0].font.color.rgb = RGBColor(31, 78, 121)

            cell_cont = row.cells
            # Si contiene saltos o viñetas, separarlos limpiamente
            lineas = contenido.replace(";", "\n").split("\n")
            cell_cont.text = ""
            for idx_l, l in enumerate(lineas):
                if l.strip():
                    p = cell_cont.add_paragraph() if idx_l > 0 else cell_cont.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.add_run(l.strip()).font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Guardar archivo
    nombre_limpio = "".join(c for c in paciente.nombre if c.isalnum() or c in (' ', '_')).strip().replace(' ', '_')
    nombre_archivo = f"PAE_{nombre_limpio}_{paciente.id}.docx"
    ruta_completa = os.path.join(CARPETA_REPORTES, nombre_archivo)
    doc.save(ruta_completa)

    return ruta_completa